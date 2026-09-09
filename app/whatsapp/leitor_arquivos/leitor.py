import csv
import json
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


EXTENSOES_SUPORTADAS = {
    ".txt",
    ".csv",
    ".json",
    ".xlsx",
    ".pdf",
}


def ler_arquivo(caminho):

    caminho = Path(caminho)

    if not caminho.exists():

        return {
            "ok": False,
            "erro": "Arquivo não encontrado."
        }

    extensao = caminho.suffix.lower()

    if extensao not in EXTENSOES_SUPORTADAS:

        return {
            "ok": False,
            "erro": (
                f"Formato "
                f"{extensao or 'sem extensão'} "
                f"ainda não suportado."
            )
        }

    try:

        if extensao == ".txt":
            return ler_txt(caminho)

        if extensao == ".csv":
            return ler_csv(caminho)

        if extensao == ".json":
            return ler_json(caminho)

        if extensao == ".xlsx":
            return ler_xlsx(caminho)

        if extensao == ".pdf":
            return ler_pdf(caminho)

        return {
            "ok": False,
            "erro": "Formato não suportado."
        }

    except Exception as exc:

        print(
            "ERRO AO LER ARQUIVO:",
            exc
        )

        return {
            "ok": False,
            "arquivo": caminho.name,
            "extensao": extensao,
            "erro": str(exc)
        }


def ler_txt(caminho):

    texto = caminho.read_text(
        encoding="utf-8-sig",
        errors="replace"
    )

    return {
        "ok": True,
        "arquivo": caminho.name,
        "extensao": ".txt",
        "texto": texto.strip()
    }


def ler_csv(caminho):

    linhas = []

    with caminho.open(
        "r",
        encoding="utf-8-sig",
        newline=""
    ) as arquivo:

        leitor = csv.reader(arquivo)

        for linha in leitor:

            linhas.append(linha)

    texto_linhas = []

    for linha in linhas:

        texto_linhas.append(
            " | ".join(
                str(item).strip()
                for item in linha
            )
        )

    return {
        "ok": True,
        "arquivo": caminho.name,
        "extensao": ".csv",
        "texto": (
            "\n".join(texto_linhas)
            .strip()
        ),
        "linhas": linhas
    }


def ler_json(caminho):

    with caminho.open(
        "r",
        encoding="utf-8-sig"
    ) as arquivo:

        dados = json.load(arquivo)

    texto = json.dumps(
        dados,
        ensure_ascii=False,
        indent=2
    )

    return {
        "ok": True,
        "arquivo": caminho.name,
        "extensao": ".json",
        "texto": texto,
        "dados": dados
    }


def ler_xlsx(caminho):

    workbook = load_workbook(
        filename=caminho,
        read_only=True,
        data_only=True
    )

    abas = {}
    texto_abas = []

    try:

        for nome_aba in workbook.sheetnames:

            planilha = workbook[nome_aba]

            linhas = []

            for linha in planilha.iter_rows(
                values_only=True
            ):

                valores = []

                for valor in linha:

                    if valor is None:

                        valores.append("")

                    else:

                        valores.append(
                            str(valor).strip()
                        )

                if any(valores):

                    linhas.append(valores)

            abas[nome_aba] = linhas

            texto_abas.append(
                f"[Aba: {nome_aba}]"
            )

            for linha in linhas:

                texto_abas.append(
                    " | ".join(linha)
                )

    finally:

        workbook.close()

    return {
        "ok": True,
        "arquivo": caminho.name,
        "extensao": ".xlsx",
        "texto": (
            "\n".join(texto_abas)
            .strip()
        ),
        "abas": abas
    }


def ler_pdf(caminho):

    leitor = PdfReader(str(caminho))

    paginas = []
    texto_paginas = []

    for numero, pagina in enumerate(
        leitor.pages,
        start=1
    ):

        texto = pagina.extract_text() or ""

        texto = texto.strip()

        paginas.append({
            "pagina": numero,
            "texto": texto
        })

        if texto:

            texto_paginas.append(
                f"[Página {numero}]"
            )

            texto_paginas.append(
                texto
            )

    texto_final = (
        "\n\n".join(texto_paginas)
        .strip()
    )

    return {
        "ok": True,
        "arquivo": caminho.name,
        "extensao": ".pdf",
        "texto": texto_final,
        "paginas": paginas,
        "quantidade_paginas": len(
            leitor.pages
        )
    }


def arquivo_suportado(caminho):

    caminho = Path(caminho)

    return (
        caminho.suffix.lower()
        in EXTENSOES_SUPORTADAS
    )


def resumir_arquivo(caminho):

    resultado = ler_arquivo(caminho)

    if not resultado.get("ok"):

        return resultado

    texto = resultado.get(
        "texto",
        ""
    )

    return {
        "ok": True,
        "arquivo": resultado.get(
            "arquivo"
        ),
        "extensao": resultado.get(
            "extensao"
        ),
        "tamanho_texto": len(texto),
        "texto": texto
    }
